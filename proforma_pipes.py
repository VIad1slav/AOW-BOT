#!/usr/bin/env python3
"""Specification XLSX → Faktura Pro Forma (DOCX) for the GOLFSTREAM line.

This is AOW's second line of documents — sanitary ware (Skolan, KG, HT), some
forty positions per invoice, always the same buyer and always the same delivery
terms. Unlike the car line, nothing here has to be typed in by hand except the
Pro Forma number and its date: the specification already carries the articles,
the descriptions, the quantities and the prices.

Two traps live in the source file, and both are silent:

* the sheet holds TWO price columns — F/G is what AOW sells at, I/J is what it
  buys at. Reading the wrong pair produces a perfectly plausible invoice with
  the wrong money in it, so the columns are located by their headers and the
  leftmost (resale) pair is the one taken;
* the totals are formulas. openpyxl serves the cached results Excel left behind,
  but only when Excel itself wrote the file — so the totals are recomputed here
  and the cached ones are used to check that arithmetic, never to replace it.
"""

import re
from copy import deepcopy
from pathlib import Path

from docx import Document

from docx_text import doc_tables, iter_paragraphs, replace_in_paragraph
from proforma import fmt_money

BASE_DIR = Path(__file__).parent
TEMPLATE = BASE_DIR / 'templates' / 'proforma_pipes_template.docx'

# Payment runs long on this line: the goods are made to order, so the deadline
# is 25 days after the invoice date (the car line uses 7).
TERMIN_DAYS = 25


# -- Money and quantity -------------------------------------------------------
def fmt_amount(value) -> str:
    """3609.6 → '3609,60'. Item rows carry no thousands separator; only the
    Razem row and the summary lines below the table do."""
    return '{:.2f}'.format(float(value)).replace('.', ',')


def fmt_qty(value) -> str:
    """11135 → '11 135' for the Razem row. Item quantities are plain integers."""
    return '{:,d}'.format(int(round(float(value)))).replace(',', ' ')


def _num(cell_value):
    """A number out of a spreadsheet cell, whatever it was typed as."""
    if cell_value is None:
        return None
    if isinstance(cell_value, (int, float)):
        return float(cell_value)
    s = str(cell_value).strip().replace(' ', '').replace('\xa0', '')
    s = s.replace(',', '.')
    if not re.fullmatch(r'-?\d*\.?\d+', s or ''):
        return None
    return float(s)


# -- The reference line at the foot of the invoice ----------------------------
def refs_from_name(name):
    """'Proforma AOW - Golfstream - Specification 126 (AB01139034).xlsx'
    → ('126', '1139034').

    The code in brackets is written with a letter prefix and a leading zero
    ("AB01139034"); the invoice shows neither. Anything missing comes back as an
    empty string — it is a suggestion the user confirms, not a fact."""
    stem = Path(name or '').stem
    spec = ''
    m = re.search(r'specification\s*[N№#]?\s*(\d+)', stem, re.IGNORECASE)
    if not m:
        m = re.search(r'spec[a-z]*[_\s№#]+(\d+)', stem, re.IGNORECASE)
    if m:
        spec = m.group(1)

    code = ''
    for chunk in re.findall(r'\(([^)]*)\)', stem):
        digits = re.sub(r'\D', '', chunk)
        if digits:
            code = digits.lstrip('0') or '0'
            break
    return spec, code


# -- Reading the specification ------------------------------------------------
# The header row is found by its labels, so a title row above the table (or a
# column shifted sideways) does not throw the reader off.
_COLS = {
    'lp':      ('no', '№', 'nr'),
    'artikel': ('item', 'artikel', 'art', 'артикул'),
    'name':    ('description', 'name'),
    'qty':     ('quantity', 'qty', 'ilosc', 'ilość', 'количество'),
    'price':   ('unit price', 'cena', 'цена'),
    'amount':  ('total eur', 'total', 'amount', 'сумма'),
}

_PRICE_LABELS = _COLS['price']
_AMOUNT_LABELS = _COLS['amount']


def _label(cell_value) -> str:
    return re.sub(r'\s+', ' ', str(cell_value or '')).strip().lower()


def _find_header(ws):
    """Row number and column map of the goods table's header."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        labels = {}
        for c in row:
            text = _label(c.value)
            if text and text not in labels:
                labels[text] = c.column
        if not labels:
            continue
        hit = {}
        for field, names in _COLS.items():
            for want in names:
                if want in labels:
                    hit[field] = labels[want]
                    break
        # "Item" and "Quantity" together only ever appear on the goods header.
        if 'artikel' in hit and 'qty' in hit:
            return row[0].row, hit
    raise ValueError('в спецификации не найдена шапка таблицы товара '
                     '(нужны колонки Item и Quantity)')


def _price_columns(ws, header_row):
    """The resale price pair.

    The sheet repeats "Unit Price" / "Total EUR" twice: the first pair is what
    AOW sells at, the second what it buys at. Both are collected here and the
    leftmost of each is taken."""
    prices, totals = [], []
    for cell in ws[header_row]:
        text = _label(cell.value)
        if text in _PRICE_LABELS:
            prices.append(cell.column)
        elif text in _AMOUNT_LABELS:
            totals.append(cell.column)
    return (min(prices) if prices else None,
            min(totals) if totals else None)


def read_spec(path, log=None):
    """Read the specification. Returns
    {'items': [...], 'qty': float, 'total': float, 'sheet': str}."""
    import openpyxl

    def note(m):
        if log:
            log(m)

    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.worksheets[0]
    header_row, cols = _find_header(ws)
    price_col, amount_col = _price_columns(ws, header_row)
    if not price_col:
        raise ValueError('в спецификации нет колонки с ценой (Unit Price)')
    name_col = cols.get('name') or cols['artikel']
    note('📊 Лист «{}», шапка в строке {}, цена — колонка {}'.format(
        ws.title, header_row, price_col))

    items, cached_total, cached_qty = [], None, None
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        def val(col):
            return row[col - 1].value if col and col <= len(row) else None

        first = str(val(cols.get('lp') or 1) or '').strip()
        # The description goes into the invoice exactly as it stands in the
        # sheet, trailing space and all — that is what a hand-made Pro Forma
        # on this line contains, and the two documents have to match.
        name_raw = str(val(name_col) or '')
        name = name_raw.strip()
        qty = _num(val(cols['qty']))

        # The TOTAL line closes the table. Its cached numbers are the
        # independent check on our own arithmetic, so they are picked up here.
        if name.lower() == 'total' or first.lower() == 'total':
            cached_qty = qty
            cached_total = _num(val(amount_col))
            break
        if not first and not name:
            continue

        artikel = str(val(cols['artikel']) or '').strip()
        price = _num(val(price_col))
        if qty is None or price is None:
            raise ValueError(
                'строка {} ({}): не прочитаны количество или цена'.format(
                    row[0].row, artikel or name or '?'))
        if qty <= 0 or price <= 0:
            raise ValueError(
                'строка {} ({}): количество или цена не больше нуля'.format(
                    row[0].row, artikel or name or '?'))
        items.append({
            'artikel': artikel,
            'name': name_raw,
            'qty': qty,
            'price': price,
            'amount': round(qty * price, 2),
        })

    if not items:
        raise ValueError('в спецификации нет ни одной позиции')

    qty_total = sum(i['qty'] for i in items)
    total = round(sum(i['amount'] for i in items), 2)

    # Excel's own results, when it left them behind. A mismatch means the sheet
    # is not what it looks like — better to stop than to invoice the wrong sum.
    if cached_total is not None and abs(cached_total - total) > 0.01:
        raise ValueError(
            'итог не сходится: в спецификации {:.2f}, по строкам {:.2f}'.format(
                cached_total, total))
    if cached_qty is not None and abs(cached_qty - qty_total) > 0.5:
        raise ValueError(
            'количество не сходится: в спецификации {:g}, по строкам {:g}'.format(
                cached_qty, qty_total))
    if cached_total is None:
        note('⚠️ В спецификации нет посчитанных итогов — сверить сумму не с чем')

    note('✅ Позиций {}, штук {}, на сумму {} EUR'.format(
        len(items), fmt_qty(qty_total), fmt_money(total)))
    return {'items': items, 'qty': qty_total, 'total': total, 'sheet': ws.title}


# -- Writing the Pro Forma ----------------------------------------------------
def _fill(paragraphs, mapping):
    for par in paragraphs:
        if '{{' not in par.text:
            continue
        for key, value in mapping.items():
            replace_in_paragraph(par, '{{%s}}' % key, value)


def _row_paragraphs(row):
    for cell in row.cells:
        for par in cell.paragraphs:
            yield par


def build_pipes_proforma(out_path, ctx, template=None, log=None):
    """Write the Pro Forma. `ctx` needs: pf_num, date, termin, spec, code and a
    non-empty list of items as read_spec returns them."""
    def note(m):
        if log:
            log(m)

    items = ctx['items']
    if not items:
        raise ValueError('нет ни одной позиции для Pro Forma')

    doc = Document(str(template or TEMPLATE))
    tables = doc_tables(doc)
    if not tables:
        raise ValueError('в шаблоне нет таблицы с товаром')
    goods = tables[0]

    ITEM = 1                       # row 0 is the header, row 1 the sample item
    for _ in range(len(items) - 1):
        goods.rows[ITEM]._tr.addnext(deepcopy(goods.rows[ITEM]._tr))

    total = 0.0
    qty_total = 0.0
    for i, it in enumerate(items):
        amount = round(float(it['qty']) * float(it['price']), 2)
        total += amount
        qty_total += float(it['qty'])
        _fill(_row_paragraphs(goods.rows[ITEM + i]), {
            'LP': str(i + 1),
            'ARTIKEL': it['artikel'],
            'NAZWA': it['name'] or '',
            'QTY': '{:g}'.format(float(it['qty'])),
            'CENA': fmt_amount(it['price']),
            'WARTOSC': fmt_amount(amount),
            'BRUTTO': fmt_amount(amount),
        })
    total = round(total, 2)

    # Everything left: the header block, the Razem row, the summary lines and
    # the specification reference at the foot of the page.
    _fill(iter_paragraphs(doc), {
        'PF_NUM': ctx['pf_num'],
        'DATA': ctx['date'],
        'TERMIN': ctx['termin'],
        'RAZEM_QTY': fmt_qty(qty_total),
        'RAZEM': fmt_money(total),
        'SPEC': ctx.get('spec', ''),
        'CODE': ctx.get('code', ''),
    })

    leftover = sorted({name for par in iter_paragraphs(doc)
                       for name in re.findall(r'\{\{(\w+)\}\}', par.text)})
    if leftover:
        raise ValueError('в шаблоне остались незаполненные поля: ' + ', '.join(leftover))

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    note('✅ Pro Forma {}: позиций {}, штук {}, итого {} EUR'.format(
        ctx['pf_num'], len(items), fmt_qty(qty_total), fmt_money(total)))
    return {'path': out_path, 'total': total,
            'count': len(items), 'qty': qty_total}


def out_filename(pf_num) -> str:
    """The GOLFSTREAM line spells it 'Faktura Pro Forma', with a capital F —
    the car line uses 'Faktura Pro forma'. Both are kept as they are."""
    safe = re.sub(r'[\\/:*?"<>|]', '-', pf_num or '').strip()
    return 'Faktura Pro Forma {}.docx'.format(safe)
