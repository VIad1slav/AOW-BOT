#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""CMR consignment note for the GOLFSTREAM line.

The goods block of a CMR repeats what the specification's footer already says:
one line per HS code with its packing description and gross weight, then a
total. So the note is built from the finished specifications rather than asked
for again — only the goods name per HS code lives outside them, in cmr_goods.json.

One truck gets one CMR even when the load was invoiced as several
specifications ("догрузы"), so the footers of all of them are merged: one line
per HS code, weights added up, packing descriptions joined.
"""

import json
import os
import re

TEMPLATE = os.path.join(os.path.dirname(__file__), 'templates', 'cmr_sk_template.docx')
GOODS_FILE = os.path.join(os.path.dirname(__file__), 'cmr_goods.json')

# Table layout of the template, verified against the customer's own sample.
DATE_CELLS = [(19, 3), (44, 16)]
DOCS_CELL = (21, 0)
GOODS_ROWS = list(range(24, 30))     # six lines fit on the form
TOTAL_ROW = 30
COL_DESC, COL_HS, COL_WEIGHT = 0, 24, 29


def load_goods_names():
    """HS code → short goods name, as written on the customer's CMR."""
    try:
        with open(GOODS_FILE, encoding='utf-8') as f:
            return json.load(f)
    except (FileNotFoundError, ValueError):
        return {}


def _fmt_weight(value):
    """11225.9 → '11 225,90', the way the sample CMR writes it."""
    if not isinstance(value, (int, float)):
        return ''
    return f'{value:,.2f}'.replace(',', ' ').replace('.', ',')


def _fmt_places(text, goods_name):
    """Packing description for the CMR, plus the goods name.

    The wording is carried over from the specification exactly as written —
    "50 crates; 8 pipes (1 pallet); 11 cartons (2 pallets)". Customs checks the
    waybill against the specification, so rewriting the separators there only
    creates differences between two documents that are supposed to agree.
    """
    joined = re.sub(r'\s{2,}', ' ', str(text or '')).strip().rstrip(';').strip()
    return f'{joined} {goods_name}'.strip() if goods_name else joined


# Итог мест поставщик подписывает двумя разными способами. Либо это подпись
# «TOTAL COLLI:» с числом в соседней ячейке, либо целая фраза в одной ячейке —
# «Total colli: 754 packages/107 places = 106 pallets&crates + 1 carton», то
# есть уже готовая строка накладной. Второй вид разбирается этим выражением:
# отдельной ячейки с числом там нет, и без разбора итог выходил пустым.
_TOTAL_LINE = re.compile(
    r'total\s+colli+\s*:\s*'
    r'(?:(\d+)\s*packages?\s*/\s*)?'
    r'(\d+)\s*places?'
    r'(?:\s*=\s*(.+?))?\s*$', re.IGNORECASE)


def _label(ws, row):
    """The footer label in column C, normalised.

    Suppliers are not consistent: "TOTAL GROSS WEIGHT:" in one specification is
    "TOTAL GROSS WEIGHT, KG:" in the next, and "TOTAL COLLI" is sometimes
    spelled "TOTAL COLLII". Comparing raw text silently loses whole groups.
    """
    v = ws.cell(row, 3).value
    if not isinstance(v, str):
        return ''
    s = v.strip().lower().rstrip(':').strip()
    return re.sub(r',\s*kg$', '', s).strip()


def groups_from_specs(spec_paths, log=lambda m: None):
    """Read the footers of the finished specifications and merge them.

    Returns (groups, total_weight, (packs, colli, note)) where a group is
    {'hs', 'places', 'gross'}. Weights are formulas in the file, so they are
    evaluated the same way the column widths are.
    """
    from openpyxl import load_workbook
    from processing import _formula_values

    merged, order = {}, []
    packs = colli = 0.0
    notes = []

    for path in spec_paths:
        ws = load_workbook(path).active
        value_of = _formula_values(ws)

        def text(r, c):
            v = ws.cell(r, c).value
            return str(v).strip() if isinstance(v, str) else ('' if v is None else str(v))

        hs_rows = [r for r in range(1, ws.max_row + 1)
                   if _label(ws, r) == 'harmonized system code']

        for i, hr in enumerate(hs_rows):
            end = hs_rows[i + 1] if i + 1 < len(hs_rows) else ws.max_row + 1
            places, gross = '', None
            for r in range(hr, end):
                lab = _label(ws, r)
                # The FIRST weight row inside a group is the group's own. The
                # last group's range also reaches the sheet's grand total, and
                # taking that one silently reported the whole truck's weight as
                # if it belonged to one HS code.
                # Одни спецификации подписывают эту строку «places:»,
                # другие «Colli:». Грандиозный итог ниже называется
                # «TOTAL COLLI» и сюда не попадает — он длиннее.
                if lab in ('places', 'colli') and not places:
                    places = text(r, 4)
                elif lab.startswith('total gross weight') and gross is None:
                    gross = value_of(r, 4)

            hs = text(hr, 4)
            if hs in merged:
                if places:
                    merged[hs]['places'] += '; ' + places
                if isinstance(gross, (int, float)):
                    merged[hs]['gross'] = (merged[hs]['gross'] or 0) + gross
            else:
                merged[hs] = {'hs': hs, 'places': places, 'gross': gross}
                order.append(hs)

        # The grand totals sit below every group, so the last row wins: in a
        # long specification each group carries its own "Total colli" too.
        spec_colli = spec_packs = None
        spec_note = ''
        for r in range(1, ws.max_row + 1):
            lab = _label(ws, r)
            if re.fullmatch(r'total colli+', lab):
                v = value_of(r, 4)
                if isinstance(v, (int, float)):
                    spec_colli = v
                # Расшифровка по видам тары из соседней ячейки в накладную не
                # идёт — решено по образцу CMR SK 114.
            elif lab == 'total packs':
                v = value_of(r, 4)
                if isinstance(v, (int, float)):
                    spec_packs = v
            elif lab.startswith('total colli'):
                m = _TOTAL_LINE.match(text(r, 3))
                if m:
                    if m.group(1):
                        spec_packs = int(m.group(1))
                    spec_colli = int(m.group(2))
                    # Хвост здесь — часть фразы, которую написал поставщик, а
                    # не наша расшифровка, поэтому он переносится дословно.
                    spec_note = (m.group(3) or '').strip()
                else:
                    log('  ⚠ CMR: не разобрана строка итога — ' + text(r, 3))
        colli += spec_colli or 0
        packs += spec_packs or 0
        if spec_note:
            notes.append(spec_note)

    groups = [merged[hs] for hs in order]
    total_weight = sum(g['gross'] for g in groups if isinstance(g['gross'], (int, float)))

    missing = [g['hs'] for g in groups if not isinstance(g['gross'], (int, float))]
    if missing:
        log('  ⚠ CMR: не удалось посчитать вес для кода ' + ', '.join(missing))

    log(f'  ✓ CMR: спецификаций {len(spec_paths)}, строк груза {len(groups)}, '
        f'общий вес {_fmt_weight(total_weight)}')
    return groups, total_weight, (packs or None, colli or None, '; '.join(notes))


def _set_cell(cell, value):
    """Write text into a template cell, keeping the font it was set up with."""
    paragraphs = cell.paragraphs
    if not paragraphs:
        return
    p = paragraphs[0]
    if p.runs:
        p.runs[0].text = str(value)
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(str(value))
    for extra in paragraphs[1:]:
        for r in extra.runs:
            r.text = ''


def invoice_short(invoice_num):
    """'FV26-111' → '111' — the number the file is named after."""
    s = re.sub(r'^FV', '', str(invoice_num or '').strip(), flags=re.IGNORECASE)
    return s.rsplit('-', 1)[-1] if '-' in s else s


def build_cmr(spec_paths, params, dst, log=lambda m: None):
    """Fill the CMR template from the finished specifications."""
    from docx import Document

    groups, total_weight, (packs, colli, note) = groups_from_specs(spec_paths, log)
    names = load_goods_names()

    doc = Document(TEMPLATE)
    tbl = doc.tables[0]

    date_dots = str(params.get('date', ''))
    date_slash = date_dots.replace('.', '/')
    for r, c in DATE_CELLS:
        _set_cell(tbl.rows[r].cells[c], date_slash)

    _set_cell(tbl.rows[DOCS_CELL[0]].cells[DOCS_CELL[1]],
              f"Invoice № {params.get('invoice_num', '')} dtd {date_dots}")

    if len(groups) > len(GOODS_ROWS):
        log(f'  ⚠ CMR: строк груза {len(groups)}, а на бланке помещается '
            f'{len(GOODS_ROWS)} — лишние не поместились')

    unknown = []
    for row_idx, group in zip(GOODS_ROWS, groups):
        name = names.get(str(group['hs']), '')
        if not name:
            unknown.append(str(group['hs']))
        row = tbl.rows[row_idx]
        _set_cell(row.cells[COL_DESC], _fmt_places(group['places'], name))
        _set_cell(row.cells[COL_HS], group['hs'])
        _set_cell(row.cells[COL_WEIGHT], _fmt_weight(group['gross']))
    if unknown:
        log('  ⚠ CMR: нет названия груза для кода ' + ', '.join(unknown) +
            ' — допишите в cmr_goods.json')

    bits = []
    if isinstance(packs, (int, float)):
        bits.append(f'{packs:g} packages')
    if isinstance(colli, (int, float)):
        bits.append(f'{colli:g} places')
    line = 'Total colli: ' + '/'.join(bits) if bits else 'Total colli:'
    # Хвост дописывается, только когда спецификация сама написала итог фразой:
    # накладная обязана совпадать со спецификацией слово в слово. Собственную
    # расшифровку по видам тары бот не сочиняет — решено по образцу CMR SK 114.
    if note:
        line += f' = {note}'
    _set_cell(tbl.rows[TOTAL_ROW].cells[COL_DESC], line)
    _set_cell(tbl.rows[TOTAL_ROW].cells[COL_WEIGHT], _fmt_weight(total_weight))

    doc.save(dst)
    log(f'  💾 CMR сохранён: {os.path.basename(dst)}')
    return dst
