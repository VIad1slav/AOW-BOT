#!/usr/bin/env python3
"""Run-aware text surgery on a .docx.

Word splits a single visible sentence into many <w:r> runs (spell-check marks,
language runs, an edit made years ago). "PF26-8/106" in the sample Pro Forma is
seven separate runs. So every replacement here works on the paragraph's joined
text and then writes the result back across the runs it covered, keeping the
formatting of the run the match started in.
"""

from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_paragraphs(doc):
    """Every paragraph in the document, including those inside tables."""
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag.endswith('}p'):
            yield Paragraph(child, doc)
        elif child.tag.endswith('}tbl'):
            for par in iter_table_paragraphs(Table(child, doc)):
                yield par


def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                yield par
            for inner in cell.tables:
                for par in iter_table_paragraphs(inner):
                    yield par


def replace_in_paragraph(par, old, new, limit=0):
    """Replace `old` with `new` inside one paragraph, across run boundaries.

    `new` may contain \n — python-docx turns those into real <w:br/> line
    breaks, which is how the buyer block and the goods cell keep their layout.
    Returns the number of replacements made.
    """
    if not old:
        return 0
    done = 0
    # Searching resumes past what was just written, so a `new` that contains
    # `old` ("EUR" → "EUR EUR") replaces each occurrence once instead of
    # matching its own output forever.
    start = 0
    while True:
        runs = par.runs
        if not runs:
            break
        texts = [r.text for r in runs]
        idx = ''.join(texts).find(old, start)
        if idx < 0:
            break
        end = idx + len(old)
        start = idx + len(new)

        spans, pos = [], 0
        for t in texts:
            spans.append((pos, pos + len(t)))
            pos += len(t)

        first = None
        for i, (s, e) in enumerate(spans):
            if s <= idx < e or (s == e == idx):
                first = i
                break
        if first is None:                     # match starts past the last run
            break

        for i in range(first, len(runs)):
            s, e = spans[i]
            if s >= end and i != first:
                break
            before = texts[i][:idx - s] if s < idx else ''
            after = texts[i][end - s:] if e > end else ''
            runs[i].text = (before + new + after) if i == first else after

        done += 1
        if limit and done >= limit:
            break
    return done


def replace_everywhere(doc, old, new, limit=0):
    total = 0
    for par in iter_paragraphs(doc):
        total += replace_in_paragraph(par, old, new, limit)
        if limit and total >= limit:
            break
    return total


def replace_in_cell(table, row, col, old, new):
    cell = table.rows[row].cells[col]
    total = 0
    for par in cell.paragraphs:
        total += replace_in_paragraph(par, old, new)
    return total


def doc_tables(doc):
    return [Table(c, doc) for c in doc.element.body.iterchildren() if c.tag.endswith('}tbl')]
