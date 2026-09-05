#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Checks for the CMR consignment note.

The note is not typed in — its goods block is lifted out of the finished
specification. So the parts worth checking are the seam between the two: does
the footer come across correctly, is the packing wording rewritten the way the
customer writes it, and does the form still carry its constant data afterwards.
"""

import os
import sys
import tempfile
import unicodedata

import cmr

SPEC = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    'samples', 'spec_golfstream.xlsx')

failures = []


def _norm(v):
    """Word stores accented letters decomposed: the "á" of "Geologická" is
    an "a" followed by a separate accent mark. Compare in one normal form so a
    correct document doesn't fail on an invisible difference."""
    return unicodedata.normalize('NFC', v) if isinstance(v, str) else v


def check(name, got, want):
    got, want = _norm(got), _norm(want)
    ok = got == want
    print(('  ✅ ' if ok else '  ❌ ') + name)
    if not ok:
        print(f'       получили: {got!r}')
        print(f'       ожидали:  {want!r}')
        failures.append(name)


def test_places_wording():
    print('\nЗапись мест: из спецификации в накладную')
    check('запись переносится дословно',
          cmr._fmt_places('27 crates; 8 cartons (1 pallet)', ''),
          '27 crates; 8 cartons (1 pallet)')
    check('название груза дописывается в конец',
          cmr._fmt_places('9 cartons (1 pallet)', 'Sealtape'),
          '9 cartons (1 pallet) Sealtape')
    check('скобка без числа остаётся как есть',
          cmr._fmt_places('4 cartons (part of pallets)', 'Blade'),
          '4 cartons (part of pallets) Blade')
    check('двойные пробелы схлопываются',
          cmr._fmt_places('1 crate;  2 cartons', ''), '1 crate; 2 cartons')
    check('пустая строка не ломает', cmr._fmt_places('', ''), '')


def test_weight_format():
    print('\nФормат веса, как на бланке')
    check('тысячи через пробел, дробь через запятую',
          cmr._fmt_weight(11225.9), '11 225,90')
    check('меньше тысячи', cmr._fmt_weight(109.2), '109,20')
    check('копейки не теряются', cmr._fmt_weight(0.49), '0,49')
    check('не число — пусто', cmr._fmt_weight(None), '')


def test_file_name():
    print('\nИмя файла: номер инвойса без префикса')
    check('FV26-111', cmr.invoice_short('FV26-111'), '111')
    check('FV26-99', cmr.invoice_short('FV26-99'), '99')
    check('другой год', cmr.invoice_short('FV27-5'), '5')
    check('уже без FV', cmr.invoice_short('26-111'), '111')


def test_goods_names():
    print('\nСправочник названий груза')
    names = cmr.load_goods_names()
    check('код труб', names.get('39172290'), 'PP')
    check('код фитингов', names.get('39174000'), 'Fittings')
    check('незнакомый код — пусто', names.get('00000000', ''), '')


def test_from_specification():
    print('\nГрузовой блок берётся из готовой спецификации')
    if not os.path.exists(SPEC):
        check('образец спецификации на месте', False, True)
        return
    groups, total, (packs, colli, note) = cmr.groups_from_specs([SPEC])
    check('групп груза', len(groups), 2)
    check('код первой группы', groups[0]['hs'], '39172290')
    check('код второй группы', groups[1]['hs'], '39174000')
    # Вес группы — её собственный итог, а не итог всей спецификации. Диапазон
    # последней группы дотягивается до общей строки внизу листа, и если брать
    # последнее совпадение, вес всей машины попадает в одну строку груза.
    check('вес первой группы', round(groups[0]['gross'], 2), 6547.29)
    check('вес второй группы', round(groups[1]['gross'], 2), 4685.92)
    check('общий вес — сумма групп', round(total, 2), 11233.21)
    check('сумма сходится с итогом спецификации',
          round(groups[0]['gross'] + groups[1]['gross'], 2), 11233.21)
    check('места переведены на английский', 'пал' not in groups[0]['places'], True)


def test_filled_form():
    print('\nЗаполненный бланк')
    if not os.path.exists(SPEC):
        check('образец спецификации на месте', False, True)
        return
    from docx import Document
    params = {'invoice_num': 'FV26-111', 'date': '29.08.2026'}
    with tempfile.TemporaryDirectory() as tmp:
        dst = os.path.join(tmp, 'CMR SK 111.docx')
        cmr.build_cmr([SPEC], params, dst)
        t = Document(dst).tables[0]

        check('дата приёма груза', t.rows[19].cells[3].text.strip(), '29/08/2026')
        check('дата составления', t.rows[44].cells[16].text.strip(), '29/08/2026')
        check('ссылка на инвойс', t.rows[21].cells[0].text.strip(),
              'Invoice № FV26-111 dtd 29.08.2026')
        check('первая строка груза', t.rows[24].cells[0].text.strip(),
              '27 crates; 8 cartons (1 pallet); 1 transport box (1 pallet) PP')
        check('её код ТН ВЭД', t.rows[24].cells[24].text.strip(), '39172290')
        check('её вес', t.rows[24].cells[29].text.strip(), '6 547,29')
        check('вторая строка названа Fittings',
              t.rows[25].cells[0].text.strip().endswith('Fittings'), True)
        check('итоговый вес', t.rows[30].cells[29].text.strip(), '11 233,21')
        check('итоговая строка начинается как надо',
              t.rows[30].cells[0].text.strip().startswith('Total colli:'), True)
        check('лишние строки груза пустые', t.rows[26].cells[0].text.strip(), '')

        print('\n  Постоянные поля бланка не затронуты')
        check('отправитель', t.rows[1].cells[0].text.strip(), 'AOW GROUP SP. z o.o.')
        check('получатель', t.rows[8].cells[0].text.strip(), 'GOLFSTREAM s.r.o.')
        check('место разгрузки', t.rows[13].cells[4].text.strip(),
              'Geologická 34, 821 06 Bratislava')
        check('условия поставки', 'FCA' in t.rows[41].cells[2].text, True)


DOGRUZ = [os.path.join(os.path.dirname(os.path.abspath(__file__)), 'samples', n)
          for n in ('spec_dogruz_1.xlsx', 'spec_dogruz_2.xlsx')]


def test_dogruzy():
    """Одна машина — одна накладная, даже если счетов было несколько."""
    print('\nДогрузы: несколько спецификаций на одну машину')
    if not all(os.path.exists(p) for p in DOGRUZ):
        check('образцы догрузов на месте', False, True)
        return
    groups, total, (packs, colli, note) = cmr.groups_from_specs(DOGRUZ)

    codes = [g['hs'] for g in groups]
    check('строка на каждый код, без повторов', len(codes), len(set(codes)))
    check('кодов получилось три', len(groups), 3)

    by_code = {g['hs']: g for g in groups}
    # 39174000 есть в обеих спецификациях: 22,25 и 3 207,35
    check('вес общего кода сложен', round(by_code['39174000']['gross'], 2), 3229.60)
    check('код только из второй', round(by_code['73269098']['gross'], 2), 74.91)
    check('общий вес по обеим', round(total, 2), 13018.40)
    check('места сложены', colli, 81)
    check('упаковки сложены', packs, 466)
    check('описания мест объединены',
          by_code['39174000']['places'].count(';') >= 1, True)

    # Ярлыки в этих файлах написаны иначе: «TOTAL GROSS WEIGHT, KG» и
    # «TOTAL COLLI» с одной «i». Раньше из-за этого терялись целые группы.
    check('вес найден несмотря на «, KG» в подписи',
          all(isinstance(g['gross'], float) for g in groups), True)


COLLI_SPEC = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                          'samples', 'spec_colli.xlsx')


def test_packaging_words():
    """Упаковку поставщик пишет по-русски, в документы она идёт по-английски."""
    print('\nПеревод упаковочных слов')
    from processing import _translate_places as tr
    check('одна труба', tr('1 труба'), '1 pipe')
    check('несколько труб', tr('8 труб'), '8 pipes')
    check('трубы среди прочего',
          tr('50 обрешеток; 8 труб (1 паллета)'), '50 crates; 8 pipes (1 pallet)')
    check('поддон', tr('4 поддона'), '4 pallets')
    check('короб не спутан с коробкой', tr('31 короб'), '31 transport boxes')
    check('коробка осталась коробкой', tr('8 коробок'), '8 cartons')


def test_footer_labels():
    """Строку мест подписывают то «places:», то «Colli:» — обе должны читаться."""
    print('\nРазные подписи в футере')
    if not os.path.exists(COLLI_SPEC):
        print('  ⏭  пропущено: нет образца с подписью Colli в samples/')
        return
    groups, total, _ = cmr.groups_from_specs([COLLI_SPEC])
    check('групп найдено', len(groups), 2)
    check('описание мест не пустое', bool(groups[0]['places']), True)
    check('это именно строка Colli',
          groups[0]['places'].startswith('50 crates'), True)
    # «TOTAL GROSS WEIGHT,KG:» без пробела перед KG тоже должно читаться
    check('вес несмотря на «,KG» без пробела',
          all(isinstance(g['gross'], float) for g in groups), True)
    check('общий вес', round(total, 2), 17780.89)


TOTAL_LINE_SPEC = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                               'samples', 'spec_total_line.xlsx')


def test_total_line_footer():
    """Итог мест бывает написан одной фразой прямо в ячейке подписи.

    В таких спецификациях нет отдельной ячейки с числом мест, и накладная
    выходила с пустой строкой «Total colli:». Ожидаемые числа взяты не из
    вывода кода, а из самой фразы спецификации и из суммы весов её групп.
    """
    print()
    print('Футер, написанный одной строкой')
    if not os.path.exists(TOTAL_LINE_SPEC):
        print('  ⏭  пропущено: нет образца с однострочным итогом в samples/')
        return
    groups, total, (packs, colli, note) = cmr.groups_from_specs([TOTAL_LINE_SPEC])
    check('упаковки взяты из фразы', packs, 754)
    check('места взяты из фразы', colli, 107)
    check('хвост фразы сохранён дословно', note, '106 pallets&crates + 1 carton')
    check('групп груза', len(groups), 3)
    check('общий вес — сумма групп', round(total, 2), 10546.50)

    from docx import Document
    params = {'invoice_num': 'FV26-69', 'date': '04.09.2026'}
    with tempfile.TemporaryDirectory() as tmp:
        dst = os.path.join(tmp, 'CMR SK 69.docx')
        cmr.build_cmr([TOTAL_LINE_SPEC], params, dst)
        t = Document(dst).tables[0]
        check('итог совпал со спецификацией слово в слово',
              t.rows[30].cells[0].text.strip(),
              'Total colli: 754 packages/107 places = 106 pallets&crates + 1 carton')
        check('итоговый вес', t.rows[30].cells[29].text.strip(), '10 546,50')


def test_total_line_parsing():
    """Разбор самой фразы — отдельно от файлов."""
    print()
    print('Разбор фразы итога')

    def parse(s):
        m = cmr._TOTAL_LINE.match(s)
        return m.groups() if m else None

    check('упаковки, места и хвост',
          parse('Total colli: 754 packages/107 places = 106 pallets&crates + 1 carton'),
          ('754', '107', '106 pallets&crates + 1 carton'))
    check('без хвоста', parse('Total colli: 1050 packages/101 places'),
          ('1050', '101', None))
    check('только места', parse('TOTAL COLLI: 81 places'), (None, '81', None))
    check('вторая «i» в подписи',
          parse('Total collii: 12 packages/3 places')[:2], ('12', '3'))
    # Голая подпись — это второй вид футера: число лежит в соседней ячейке,
    # и разбирать здесь нечего.
    check('голая подпись не разбирается', parse('Total colli'), None)


def test_template_fonts():
    """В бланке не должно остаться несуществующих имён шрифтов.

    «Times New Roman Bold» — не семейство, а начертание. Word это прощает,
    LibreOffice на сервере — нет: имя не находится, прогон уезжает в
    подстановку, и в одной строке накладной оказывались разные шрифты.
    """
    print()
    print('Шрифты бланка')
    import re as _re
    import zipfile
    if not os.path.exists(cmr.TEMPLATE):
        check('бланк на месте', False, True)
        return
    with zipfile.ZipFile(cmr.TEMPLATE) as z:
        xml = z.read('word/document.xml').decode('utf-8')
    names = set(_re.findall(r'w:(?:ascii|hAnsi|cs)="([^"]*)"', xml))
    fake = sorted(n for n in names if _re.search(r'(Bold|Italic|CYR)$', n))
    check('псевдо-имён шрифтов нет', fake, [])

    from docx import Document
    t = Document(cmr.TEMPLATE).tables[0]
    weights = [t.rows[r].cells[cmr.COL_WEIGHT].paragraphs[0].runs[0]
               for r in cmr.GOODS_ROWS]
    # Часть прогонов шрифт не называет вовсе — они берут умолчание документа.
    # Это нормально ровно до тех пор, пока умолчание тоже Times New Roman:
    # ломалось раньше не отсутствие имени, а имя несуществующего семейства.
    with zipfile.ZipFile(cmr.TEMPLATE) as z:
        styles = z.read('word/styles.xml').decode('utf-8')
    default = _re.search(r'<w:rPrDefault>.*?w:ascii="([^"]*)"', styles, _re.S)
    check('умолчание документа', default and default.group(1), 'Times New Roman')
    check('колонка веса — один шрифт на все строки',
          {r.font.name for r in weights} - {None}, {'Times New Roman'})
    check('и один размер', {r.font.size.pt for r in weights}, {9.0})
    check('жирность не потерялась', {r.font.bold for r in weights}, {True})


def test_template_logo():
    """Надпись «CMR» должна стоять по центру овала, а не на пустых местах.

    В образце она держалась на четырёх пробелах в левой части широкой рамки —
    при другом шрифте ширина пробела уезжает и буквы вылезают за овал. Проверка
    сторожит настоящее выравнивание: рамка совпадает с овалом, абзац центрован.
    """
    print()
    print('Надпись CMR в овале')
    import re as _re
    import zipfile
    if not os.path.exists(cmr.TEMPLATE):
        check('бланк на месте', False, True)
        return
    with zipfile.ZipFile(cmr.TEMPLATE) as z:
        xml = z.read('word/document.xml').decode('utf-8')

    def geometry(shape_id):
        m = _re.search(r'<v:\w+ id="%s"[^>]*style="([^"]*)"' % shape_id, xml)
        if not m:
            return None
        st = dict(p.split(':', 1) for p in m.group(1).split(';') if ':' in p)
        return tuple(st.get(k) for k in ('left', 'top', 'width', 'height'))

    oval, rect = geometry('Oval 4'), geometry('Rectangles 3')
    check('рамка надписи совпала с овалом', rect, oval)

    start = xml.index('<v:rect id="Rectangles 3"')
    box = xml[start:xml.index('</v:rect>', start)]
    check('по вертикали — по середине', 'v-text-anchor:middle' in box, True)
    check('по горизонтали — по центру', '<w:jc w:val="center"/>' in box, True)
    check('ведущих пробелов не осталось',
          _re.search(r'<w:t xml:space="preserve">\s+</w:t>', box), None)


def main():
    print('Проверка накладной CMR')
    test_packaging_words()
    test_footer_labels()
    test_total_line_footer()
    test_total_line_parsing()
    test_template_fonts()
    test_template_logo()
    test_dogruzy()
    test_places_wording()
    test_weight_format()
    test_file_name()
    test_goods_names()
    test_from_specification()
    test_filled_form()
    print('\n' + ('✅ ВСЁ ПРОШЛО' if not failures
                  else '❌ ПРОВАЛЕНО %d: %s' % (len(failures), '; '.join(failures))))
    return 1 if failures else 0


if __name__ == '__main__':
    sys.exit(main())
