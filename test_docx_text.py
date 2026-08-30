#!/usr/bin/env python3
"""Checks for the run-splicing engine in docx_text.py.

Word scatters one visible sentence across many <w:r> runs, so every
replacement has to be found in the joined paragraph text and written back
across whatever runs it happened to cover. That is fiddly enough to deserve
its own checks, built on paragraphs whose runs are deliberately fragmented.
"""

import sys

from docx import Document
from docx.shared import Pt

from docx_text import (
    doc_tables, iter_paragraphs, replace_everywhere, replace_in_cell,
    replace_in_paragraph,
)

failures = []


def check(name, got, want):
    ok = got == want
    print(('  ✅ ' if ok else '  ❌ ') + name
          + ('' if ok else '  — получено %r, ожидалось %r' % (got, want)))
    if not ok:
        failures.append(name)
    return ok


def para(*chunks):
    """A fresh paragraph whose runs are exactly `chunks`.

    A chunk is either a string, or (text, bold) when the check cares about
    which run's formatting the replacement inherits.
    """
    doc = Document()
    p = doc.add_paragraph()
    for chunk in chunks:
        text, bold = chunk if isinstance(chunk, tuple) else (chunk, None)
        run = p.add_run(text)
        run.bold = bold
    return doc, p


def fmt(p):
    """(bold-per-character) for the whole paragraph — None where unset."""
    out = []
    for run in p.runs:
        for _ in run.text:
            out.append(run.bold)
    return out


def test_basic():
    print('\nЗамена внутри абзаца')

    _, p = para('Счёт ', 'PF2', '6', '-', '8', '/', '10', '6', ' от')
    n = replace_in_paragraph(p, 'PF26-8/106', '{{NUM}}')
    check('склейка семи прогонов', (n, p.text), (1, 'Счёт {{NUM}} от'))

    _, p = para('один два три')
    check('совпадение внутри одного прогона',
          (replace_in_paragraph(p, 'два', 'ДВА'), p.text), (1, 'один ДВА три'))

    _, p = para('нача', 'ло текста')
    check('совпадение в самом начале',
          (replace_in_paragraph(p, 'начало', 'КОНЕЦ'), p.text), (1, 'КОНЕЦ текста'))

    _, p = para('текст в кон', 'це')
    check('совпадение в самом конце',
          (replace_in_paragraph(p, 'конце', 'НАЧАЛЕ'), p.text), (1, 'текст в НАЧАЛЕ'))

    _, p = para('abc')
    check('нет совпадения — 0 и текст цел',
          (replace_in_paragraph(p, 'xyz', 'zzz'), p.text), (0, 'abc'))

    _, p = para('abc')
    check('пустой образец игнорируется',
          (replace_in_paragraph(p, '', 'zzz'), p.text), (0, 'abc'))

    _, p = para('уда', 'лить это ', 'слово')
    check('замена на пустую строку удаляет',
          (replace_in_paragraph(p, 'лить это ', ''), p.text), (1, 'уда' + 'слово'))


def test_repeats():
    print('\nПовторы и ограничение')

    _, p = para('кот ', 'и ', 'кот ', 'и кот')
    check('заменяются все вхождения',
          (replace_in_paragraph(p, 'кот', 'пёс'), p.text), (3, 'пёс и пёс и пёс'))

    _, p = para('кот и кот и кот')
    check('limit=1 меняет только первое',
          (replace_in_paragraph(p, 'кот', 'пёс', limit=1), p.text),
          (1, 'пёс и кот и кот'))

    _, p = para('кот и кот и кот')
    check('limit=2 меняет два первых',
          (replace_in_paragraph(p, 'кот', 'пёс', limit=2), p.text),
          (2, 'пёс и пёс и кот'))

    # Without an advancing search offset this spins until the guard trips.
    _, p = para('EUR и EUR')
    check('замена, содержащая образец, не зацикливается',
          (replace_in_paragraph(p, 'EUR', 'EUR EUR'), p.text),
          (2, 'EUR EUR и EUR EUR'))

    _, p = para('aaaa')
    check('перекрывающиеся вхождения не считаются дважды',
          (replace_in_paragraph(p, 'aa', 'b'), p.text), (2, 'bb'))


def test_formatting():
    print('\nФорматирование')

    _, p = para(('жирное ', True), ('слово', None))
    replace_in_paragraph(p, 'жирное слово', 'ЗАМЕНА')
    check('замена берёт формат первого прогона', fmt(p), [True] * len('ЗАМЕНА'))

    _, p = para(('сумма ', True), ('5 400', True), (',00', True), (' EUR', None))
    replace_in_paragraph(p, '5 400,00', '9 999,99')
    check('хвост за совпадением сохраняет свой формат',
          (p.text, fmt(p)[-4:]), ('сумма 9 999,99 EUR', [None] * 4))

    _, p = para(('до ', None), ('X', True), (' после', None))
    replace_in_paragraph(p, 'X', 'Y')
    check('соседние прогоны не тронуты',
          (p.text, fmt(p)), ('до Y после', [None] * 3 + [True] + [None] * 6))


def test_line_breaks():
    print('\nПереносы строк')

    _, p = para('строка1')
    replace_in_paragraph(p, 'строка1', 'a\nb\nc')
    check('\\n в замене становится настоящим переносом', p.text, 'a\nb\nc')
    xml = p.runs[0]._r.xml
    check('  и это <w:br/>, а не текст', '<w:br/>' in xml, True)

    doc = Document()
    p = doc.add_paragraph()
    p.add_run('Имя')
    p.add_run().add_break()
    p.add_run('Адрес')
    p.add_run().add_break()
    p.add_run('Паспорт')
    check('исходный текст с переносами читается',
          p.text, 'Имя\nАдрес\nПаспорт')
    n = replace_in_paragraph(p, 'Адрес\nПаспорт', '{{INFO}}')
    check('совпадение через перенос строки', (n, p.text), (1, 'Имя\n{{INFO}}'))


def test_document_level():
    print('\nПо всему документу и таблицам')

    doc = Document()
    doc.add_paragraph('шапка Х')
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = 'Х в ячейке'
    table.rows[1].cells[1].text = 'ещё Х'
    doc.add_paragraph('подвал Х')

    check('iter_paragraphs обходит и таблицы',
          len([p for p in iter_paragraphs(doc) if p.text.strip()]), 4)
    check('doc_tables находит таблицу', len(doc_tables(doc)), 1)
    check('replace_everywhere меняет везде',
          replace_everywhere(doc, 'Х', 'Y'), 4)
    check('  текст обновлён',
          [p.text for p in iter_paragraphs(doc) if p.text.strip()],
          ['шапка Y', 'Y в ячейке', 'ещё Y', 'подвал Y'])

    doc = Document()
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = 'цель'
    t.rows[1].cells[1].text = 'цель'
    check('replace_in_cell бьёт только по своей ячейке',
          replace_in_cell(t, 0, 0, 'цель', 'готово'), 1)
    check('  соседняя ячейка не тронута', t.rows[1].cells[1].text, 'цель')

    doc = Document()
    doc.add_paragraph('раз').add_run(' два')
    check('replace_everywhere с limit останавливается',
          replace_everywhere(doc, 'а', 'А', limit=1), 1)


def test_survives_save():
    print('\nПереживает сохранение')
    import io
    doc = Document()
    p = doc.add_paragraph()
    for chunk in ('PF2', '6', '-', '8'):
        p.add_run(chunk)
    replace_in_paragraph(p, 'PF26-8', 'PF27-1\nвторая строка')

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    again = Document(buf)
    check('текст читается после save/open',
          again.paragraphs[0].text, 'PF27-1\nвторая строка')


def main():
    print('Проверка docx_text — замены по прогонам Word')
    test_basic()
    test_repeats()
    test_formatting()
    test_line_breaks()
    test_document_level()
    test_survives_save()
    print('\n' + ('✅ ВСЁ ПРОШЛО' if not failures
                  else '❌ ПРОВАЛЕНО %d: %s' % (len(failures), '; '.join(failures))))
    return 1 if failures else 0


if __name__ == '__main__':
    sys.exit(main())
