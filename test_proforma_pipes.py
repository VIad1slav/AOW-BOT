#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Checks for the pipes Pro Forma (the GOLFSTREAM line).

Nothing on this line is typed in, so the whole document stands or falls on the
specification being read correctly. Two things there are worth guarding above
all others, because both fail silently and both produce a document that looks
perfectly ordinary:

* the sheet carries two price columns — resale and purchase. Reading the wrong
  one costs real money and nothing on the page gives it away;
* the layout writes item amounts without a thousands separator and the totals
  with one. Getting that backwards is not an error anyone notices in review.

The strongest check is the last one: rebuild PF26-9/110 from the specification
it was made from and compare it, line by line, with the document a human made.
"""

import os
import sys
import tempfile

import openpyxl
from docx import Document

import proforma_pipes as pp

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLE_XLSX = os.path.join(HERE, 'samples', 'pipes_spec_126.xlsx')
SAMPLE_DOCX = os.path.join(HERE, 'samples', 'pipes_proforma_110.docx')

failures = []


def check(name, got, want):
    ok = got == want
    print(('  ✅ ' if ok else '  ❌ ') + name)
    if not ok:
        print(f'       получили: {got!r}')
        print(f'       ожидали:  {want!r}')
        failures.append(name)


def raises(name, fn, needle=''):
    try:
        fn()
    except Exception as e:
        if needle and needle not in str(e):
            print('  ❌ ' + name)
            print(f'       ошибка не про то: {e}')
            failures.append(name)
        else:
            print('  ✅ ' + name)
        return
    print('  ❌ ' + name)
    print('       ошибки не было, а должна была быть')
    failures.append(name)


# ── A specification built here, so the checks do not depend on a customer file ─
def make_sheet(path, rows=None, total=None, qty_total=None):
    """A sheet shaped like the real ones: a title, a header row, the goods, and
    a TOTAL line. Both price pairs are present — resale first, purchase second,
    exactly as the customer's file has them."""
    rows = rows if rows is not None else [
        (1, '116060',    'HTEM Pipe DN/OD 125х2000 mm ', 54,  6.05, 5.66),
        (2, '335040',    'Skolan Safe-EM Pipe DN/OD 110х1000 mm ', 640, 5.64, 5.27),
        (3, '770520-09', 'KG2000 EM Pipe SN10 DN/OD 160х500 mm', 35, 4.95, 4.63),
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'AOW'
    ws['B1'] = 'Proforma'
    head = ['No', 'Item', 'Наименование', ' Description', 'Quantity',
            'Unit Price', 'Total EUR', None, 'Unit Price', 'Total EUR']
    for col, value in enumerate(head, 1):
        ws.cell(4, col, value)
    r = 5
    for lp, item, name, qty, sell, buy in rows:
        ws.cell(r, 1, lp)
        ws.cell(r, 2, item)
        ws.cell(r, 3, 'русское название')
        ws.cell(r, 4, name)
        ws.cell(r, 5, qty)
        ws.cell(r, 6, sell)
        ws.cell(r, 7, round(qty * sell, 2))
        ws.cell(r, 9, buy)
        ws.cell(r, 10, round(qty * buy, 2))
        r += 1
    ws.cell(r, 4, 'TOTAL')
    ws.cell(r, 5, qty_total if qty_total is not None
            else sum(x[3] for x in rows))
    ws.cell(r, 7, total if total is not None
            else round(sum(x[3] * x[4] for x in rows), 2))
    wb.save(path)
    return path


SELL_TOTAL = round(54 * 6.05 + 640 * 5.64 + 35 * 4.95, 2)     # 4 109,55


def test_refs_from_name():
    print('\nНомер спецификации и код из имени файла')
    check('номер и код',
          pp.refs_from_name('Proforma AOW - Golfstream - Specification 126 '
                            '(AB01139034).xlsx'),
          ('126', '1139034'))
    check('буквы отбрасываются, ведущий ноль тоже',
          pp.refs_from_name('Specification 77 (XY0987654).xlsx'),
          ('77', '987654'))
    check('код без букв',
          pp.refs_from_name('Specification 5 (1139034).xlsx'),
          ('5', '1139034'))
    check('несколько нулей подряд',
          pp.refs_from_name('Specification 5 (AB00042).xlsx'), ('5', '42'))
    check('нет скобок — код пустой',
          pp.refs_from_name('Specification 12.xlsx'), ('12', ''))
    check('ничего не разобрать — обе пустые',
          pp.refs_from_name('какой-то файл.xlsx'), ('', ''))
    check('«Spec 126» тоже понимается',
          pp.refs_from_name('Spec 126 (AB01139034).xlsx'), ('126', '1139034'))


def test_number_format():
    print('\nЗапись чисел: в строках без разрядов, в итоге с разрядами')
    check('строка таблицы: без пробела', pp.fmt_amount(3609.6), '3609,60')
    check('строка таблицы: копейки всегда две', pp.fmt_amount(0.4), '0,40')
    check('строка таблицы: круглая сумма', pp.fmt_amount(1520), '1520,00')
    check('итог: разряды через пробел', pp.fmt_qty(11135), '11 135')
    check('итог: сотни без пробела', pp.fmt_qty(800), '800')
    # fmt_money comes from the car line and is shared on purpose — the summary
    # lines of both documents are written the same way.
    from proforma import fmt_money
    check('итог по деньгам: «27 827,49»', fmt_money(27827.49), '27 827,49')


def test_reads_the_resale_price():
    print('\nДве колонки цены: берётся продажная, а не закупочная')
    with tempfile.TemporaryDirectory() as tmp:
        path = make_sheet(os.path.join(tmp, 'spec.xlsx'))
        data = pp.read_spec(path)
        check('позиций', len(data['items']), 3)
        check('штук', data['qty'], 729)
        check('сумма по продажной колонке', data['total'], SELL_TOTAL)
        check('цена первой позиции', data['items'][0]['price'], 6.05)
        check('артикул', data['items'][0]['artikel'], '116060')
        check('название — английское, дословно',
              data['items'][0]['name'], 'HTEM Pipe DN/OD 125х2000 mm ')
        check('сумма позиции', data['items'][1]['amount'], round(640 * 5.64, 2))


def test_totals_are_checked_against_the_sheet():
    print('\nСверка с итогами самой спецификации')
    with tempfile.TemporaryDirectory() as tmp:
        good = make_sheet(os.path.join(tmp, 'ok.xlsx'))
        check('сходящийся итог принимается', pp.read_spec(good)['total'], SELL_TOTAL)

        bad = make_sheet(os.path.join(tmp, 'bad.xlsx'), total=9999.99)
        raises('расхождение в сумме — отказ', lambda: pp.read_spec(bad),
               'итог не сходится')

        bad_qty = make_sheet(os.path.join(tmp, 'badqty.xlsx'), qty_total=1)
        raises('расхождение в количестве — отказ', lambda: pp.read_spec(bad_qty),
               'количество не сходится')


def test_broken_sheets():
    print('\nИспорченная спецификация — отказ, а не тихая ошибка')
    with tempfile.TemporaryDirectory() as tmp:
        no_price = make_sheet(os.path.join(tmp, 'noprice.xlsx'))
        wb = openpyxl.load_workbook(no_price)
        wb.active['F6'] = None
        wb.active.cell(8, 7, None)          # итог тоже убираем, он бы не сошёлся
        wb.save(no_price)
        raises('пустая цена — отказ', lambda: pp.read_spec(no_price),
               'не прочитаны количество или цена')

        zero = make_sheet(os.path.join(tmp, 'zero.xlsx'))
        wb = openpyxl.load_workbook(zero)
        wb.active['E6'] = 0
        wb.active.cell(8, 5, None)
        wb.active.cell(8, 7, None)
        wb.save(zero)
        raises('нулевое количество — отказ', lambda: pp.read_spec(zero),
               'не больше нуля')

        empty = os.path.join(tmp, 'empty.xlsx')
        wb = openpyxl.Workbook()
        wb.active['A1'] = 'просто текст'
        wb.save(empty)
        raises('нет шапки товара — отказ', lambda: pp.read_spec(empty),
               'не найдена шапка')


def test_document():
    print('\nСобранный документ')
    with tempfile.TemporaryDirectory() as tmp:
        path = make_sheet(os.path.join(tmp, 'spec.xlsx'))
        data = pp.read_spec(path)
        out = os.path.join(tmp, pp.out_filename('PF26-9/110'))
        result = pp.build_pipes_proforma(out, {
            'pf_num': 'PF26-9/110', 'date': '2026-09-03', 'termin': '2026-09-28',
            'spec': '126', 'code': '1139034', 'items': data['items'],
        })
        check('имя файла', os.path.basename(out),
              'Faktura Pro Forma PF26-9-110.docx')
        check('итог из построителя', result['total'], SELL_TOTAL)

        doc = Document(out)
        text = '\n'.join(p.text for p in doc.paragraphs)
        check('плейсхолдеров не осталось', '{{' in text, False)
        check('номер в заголовке', 'PF26-9/110' in doc.paragraphs[0].text, True)
        check('дата выставления',
              'Warszawa, 2026-09-03' in doc.paragraphs[1].text, True)
        check('срок оплаты', '2026-09-28' in text, True)
        check('покупатель остался в шаблоне', 'GOLFSTREAM s.r.o.' in text, True)
        check('условия поставки остались',
              'FCA Warszaw Poland' in text, True)
        check('нижняя строка: номер и код',
              doc.paragraphs[-1].text.strip(), '126\t1139034')

        # The WDT clause is what lets the goods leave at all, so it is checked
        # line by line rather than by a single keyword.
        uwagi = next(p.text for p in doc.paragraphs
                     if p.text.strip().startswith('Uwagi'))
        check('заголовок WDT',
              'Wewnątrzwspólnotowa dostawa towarów (WDT)' in uwagi, True)
        for fragment in (
                '(1) [Importer/Buyer] shall not sell, export or re-export',
                'Article 12g of Council Regulation (EU) No 833/2014; and',
                'Article 8g of Council Regulation (EC) No 765/2006.',
                '(2) The [Importer/Buyer] shall use its best endeavours',
                '(3) The [Importer/Buyer] shall establish and maintain',
                '(4) Any violation of paragraphs (1), (2) or (3)',
                '- termination of this Agreement; and',
                'penalty of 10% of the total value of this Agreement',
                '(5) The [Importer/Buyer] shall promptly inform',
                'within two weeks of the simple request for such information.'):
            check('WDT: ' + fragment[:45], fragment in uwagi, True)
        check('строк в блоке Uwagi', len(uwagi.rstrip('\n').split('\n')), 13)
        check('строки «Termin dostawy» больше нет',
              'Termin dostawy' in text, False)

        table = doc.tables[0]
        check('строк: шапка + товар + итог', len(table.rows), 3 + 2)
        row = table.rows[1]
        check('LP', row.cells[0].text, '1')
        check('артикул', row.cells[1].text, '116060')
        check('название с пробелом на конце',
              row.cells[2].text, 'HTEM Pipe DN/OD 125х2000 mm ')
        check('количество целое', row.cells[3].text, '54')
        check('цена', row.cells[4].text, '6,05')
        check('сумма', row.cells[5].text, '326,70')
        check('НДС нулевой', (row.cells[6].text, row.cells[7].text), ('0', '0'))
        check('брутто = нетто', row.cells[8].text, '326,70')

        big = table.rows[2]                        # 640 × 5,64 = 3609,60
        check('в строке таблицы разрядов нет', big.cells[5].text, '3609,60')

        total_row = table.rows[-1]
        check('строка Razem', total_row.cells[2].text, 'Razem / Total')
        check('всего штук', total_row.cells[3].text, '729')
        check('всего денег', total_row.cells[5].text, '4 109,55')
        check('брутто в итоге', total_row.cells[8].text, '4 109,55')


def test_matches_the_handmade_document():
    """Rebuild PF26-9/110 from its own specification and compare with the
    document a human made from the same file."""
    print('\nСовпадение с настоящей Pro Forma PF26-9/110')
    if not (os.path.exists(SAMPLE_XLSX) and os.path.exists(SAMPLE_DOCX)):
        print('  ⏭  ПРОПУЩЕНО — нет образцов в samples/')
        return

    data = pp.read_spec(SAMPLE_XLSX)
    spec, code = pp.refs_from_name(
        'Proforma AOW - Golfstream - Specification 126 (AB01139034).xlsx')
    check('позиций', len(data['items']), 38)
    check('штук', data['qty'], 11135)
    check('сумма', data['total'], 27827.49)

    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, pp.out_filename('PF26-9/110'))
        pp.build_pipes_proforma(out, {
            'pf_num': 'PF26-9/110', 'date': '2026-09-03', 'termin': '2026-09-28',
            'spec': spec, 'code': code, 'items': data['items'],
        })

        def dump(path):
            doc = Document(path)
            lines = ['P|' + p.text for p in doc.paragraphs]
            for t in doc.tables:
                for row in t.rows:
                    lines.append('T|' + '||'.join(c.text for c in row.cells))
            return lines

        made, hand = dump(out), dump(SAMPLE_DOCX)
        check('строк столько же', len(made), len(hand))
        # The Uwagi block is the one line that is deliberately not the sample's:
        # PF26-9/110 was written before the WDT clause became mandatory, and the
        # clause is checked on its own in test_document.
        diff = [(a, b) for a, b in zip(hand, made)
                if a != b and not a.startswith('P|Uwagi')]
        check('расхождений с документом человека', len(diff), 0)
        for a, b in diff[:5]:
            print(f'       человек: {a!r}')
            print(f'       бот:     {b!r}')
        check('блок Uwagi заменён на новый',
              any(a != b and a.startswith('P|Uwagi') for a, b in zip(hand, made)),
              True)


def main():
    print('Pro Forma по спецификации — проверки')
    test_refs_from_name()
    test_number_format()
    test_reads_the_resale_price()
    test_totals_are_checked_against_the_sheet()
    test_broken_sheets()
    test_document()
    test_matches_the_handmade_document()
    print('\n' + ('✅ ВСЁ ПРОШЛО' if not failures
                  else '❌ ПРОВАЛЕНО %d: %s' % (len(failures), '; '.join(failures))))
    return 1 if failures else 0


if __name__ == '__main__':
    sys.path.insert(0, HERE)
    sys.exit(main())
